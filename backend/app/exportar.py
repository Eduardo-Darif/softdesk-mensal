"""
Exportação profissional dos chamados: Excel (várias abas) e PDF (relatório).

- Excel (openpyxl, já é dependência): aba Resumo com KPIs + quebras, abas por
  categoria/status/mês/solicitante, e a aba Chamados com o detalhe formatado.
- PDF (reportlab): um relatório enxuto pra apresentar — KPIs e as mesmas quebras
  (sem a tabela de detalhe, que num PDF viraria dezenas de páginas).

Ambos recebem uma lista de Chamado JÁ FILTRADA (o backend filtra pelos códigos
que o painel manda), então a exportação reflete exatamente o que está na tela.

O reportlab é opcional: se não estiver instalado, gerar_pdf levanta RuntimeError
com uma mensagem clara (o endpoint transforma isso num aviso pro usuário).
"""

from __future__ import annotations

import io
import re
from collections import Counter
from datetime import datetime

from .parser import Chamado

AZUL = "003366"        # azul Martinelli (cabeçalhos)
AZUL_CLARO = "DCE6F1"  # fundo suave para faixas


def _concluido(status: str) -> bool:
    s = (status or "").lower()
    s = s.translate(str.maketrans("áàâãäéêíóôõúç", "aaaaaeeiooouc"))
    return bool(re.search(r"fechad|encerrad|resolvid|cancelad|reprovad", s))


def _fmt_mes(m: str) -> str:
    nomes = ["jan", "fev", "mar", "abr", "mai", "jun",
             "jul", "ago", "set", "out", "nov", "dez"]
    try:
        ano, mes = m.split("-")
        return f"{nomes[int(mes) - 1]}/{ano[2:]}"
    except (ValueError, IndexError):
        return m


def resumir(chamados: list[Chamado]) -> dict:
    """Calcula KPIs e as quebras usadas nas duas exportações."""
    total = len(chamados)
    por_categoria = Counter((c.categoria or "(sem categoria)") for c in chamados)
    por_status = Counter((c.status or "(sem status)") for c in chamados)
    por_solicitante = Counter((c.atendente or "(sem solicitante)") for c in chamados)

    # por mês: total, concluídos e %
    meses: dict[str, dict] = {}
    for c in chamados:
        m = c.mes_abertura
        d = meses.setdefault(m, {"total": 0, "concluidos": 0})
        d["total"] += 1
        if _concluido(c.status):
            d["concluidos"] += 1
    por_mes = []
    for m in sorted(meses):
        d = meses[m]
        pct = round(d["concluidos"] / d["total"] * 100) if d["total"] else 0
        por_mes.append({"mes": m, "total": d["total"],
                        "concluidos": d["concluidos"], "pct": pct})

    concluidos = sum(1 for c in chamados if _concluido(c.status))
    return {
        "total": total,
        "concluidos": concluidos,
        "abertos": total - concluidos,
        "categoria_top": por_categoria.most_common(1)[0][0] if por_categoria else "—",
        "por_categoria": por_categoria.most_common(),
        "por_status": por_status.most_common(),
        "por_solicitante": por_solicitante.most_common(),
        "por_mes": por_mes,
    }


# ─────────────────────────── EXCEL ───────────────────────────

def gerar_xlsx(chamados: list[Chamado], descricao_filtro: str = "") -> bytes:
    """Monta o Excel em modo write_only (streaming): as linhas vão direto pro
    arquivo em vez de virar milhares de objetos de célula na memória — ~3x mais
    rápido com o ano inteiro. Nesse modo as linhas só podem ser acrescentadas
    (append), e larguras/congelamento/filtro são definidos antes delas."""
    import openpyxl
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    r = resumir(chamados)
    gerado_em = datetime.now().strftime("%d/%m/%Y %H:%M")

    hdr_fill = PatternFill("solid", fgColor=AZUL)
    hdr_font = Font(bold=True, color="FFFFFF")
    hdr_align = Alignment(horizontal="center", vertical="center")
    titulo_font = Font(bold=True, size=14, color=AZUL)
    direita = Alignment(horizontal="right")
    fina = Side(style="thin", color="D9D9D9")
    borda = Border(left=fina, right=fina, top=fina, bottom=fina)

    wb = openpyxl.Workbook(write_only=True)

    def cel(ws, valor, font=None, fill=None, alignment=None, border=None, number_format=None):
        c = WriteOnlyCell(ws, value=valor)
        if font: c.font = font
        if fill: c.fill = fill
        if alignment: c.alignment = alignment
        if border: c.border = border
        if number_format: c.number_format = number_format
        return c

    def cabecalho(ws, colunas):
        ws.append([cel(ws, nome, hdr_font, hdr_fill, hdr_align, borda) for nome in colunas])

    def largura(ws, larguras):
        for j, w in enumerate(larguras, start=1):
            ws.column_dimensions[get_column_letter(j)].width = w

    def tabela_quebra(titulo_aba, titulo, colunas, linhas, larguras, primeira_num_col=2):
        ws = wb.create_sheet(titulo_aba)
        largura(ws, larguras)
        ws.freeze_panes = "A4"
        ws.append([cel(ws, titulo, titulo_font)])
        ws.append([])
        cabecalho(ws, colunas)
        for row in linhas:
            ws.append([cel(ws, val, border=borda, alignment=direita if j >= primeira_num_col else None)
                       for j, val in enumerate(row, start=1)])

    # ── Aba Resumo ────────────────────────────────────────────
    ws = wb.create_sheet("Resumo")
    largura(ws, [34, 40])
    ws.append([cel(ws, "Relatório de Chamados — Softdesk", Font(bold=True, size=16, color=AZUL))])
    ws.append([cel(ws, "Martinelli Advogados", Font(size=11, color="808080"))])
    ws.append([f"Gerado em {gerado_em}"])
    ws.append([cel(ws, f"Filtro: {descricao_filtro or 'todos os chamados'}",
                   Font(italic=True, size=10, color="808080"))])
    ws.append([])

    kpis = [("Total de chamados", r["total"]),
            ("Concluídos", r["concluidos"]),
            ("Em aberto", r["abertos"]),
            ("Categoria mais frequente", r["categoria_top"])]
    negrito = Font(bold=True)
    for nome, val in kpis:
        ws.append([cel(ws, nome, negrito), val])

    # ── Aba Por categoria ─────────────────────────────────────
    linhas = [[k, v, f"{round(v / r['total'] * 100) if r['total'] else 0}%"]
              for k, v in r["por_categoria"]]
    tabela_quebra("Por categoria", "Chamados por categoria",
                  ["Categoria", "Qtde", "%"], linhas, [40, 12, 10])

    # ── Aba Por status ────────────────────────────────────────
    linhas = [[k, v, f"{round(v / r['total'] * 100) if r['total'] else 0}%"]
              for k, v in r["por_status"]]
    tabela_quebra("Por status", "Chamados por status",
                  ["Status", "Qtde", "%"], linhas, [30, 12, 10])

    # ── Aba Por mês ───────────────────────────────────────────
    linhas = [[_fmt_mes(d["mes"]), d["total"], d["concluidos"], f"{d['pct']}%"]
              for d in r["por_mes"]]
    tabela_quebra("Por mês", "Chamados por mês",
                  ["Mês", "Abertos no mês", "Concluídos", "% concluído"],
                  linhas, [14, 16, 14, 14])

    # ── Aba Por solicitante ───────────────────────────────────
    linhas = [[k, v] for k, v in r["por_solicitante"]]
    tabela_quebra("Por solicitante", "Chamados por solicitante",
                  ["Solicitante", "Qtde"], linhas, [40, 12])

    # ── Aba Chamados (detalhe) ────────────────────────────────
    ws = wb.create_sheet("Chamados")
    colunas = ["Código", "Título", "Categoria", "Solicitante", "Status",
               "Abertura", "Término previsto"]
    largura(ws, [12, 50, 22, 30, 20, 18, 18])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(colunas))}{max(1, len(chamados) + 1)}"
    cabecalho(ws, colunas)
    fmt_data = "DD/MM/YYYY HH:MM"
    for c in sorted(chamados, key=lambda x: x.data_abertura, reverse=True):
        ws.append([
            c.codigo, c.titulo, c.categoria, c.atendente, c.status,
            cel(ws, c.data_abertura, number_format=fmt_data),
            cel(ws, c.data_atualizacao, number_format=fmt_data) if c.data_atualizacao else None,
        ])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ─────────────────────────── PDF ───────────────────────────

def gerar_pdf(chamados: list[Chamado], descricao_filtro: str = "") -> bytes:
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import cm
        from reportlab.platypus import (Paragraph, SimpleDocTemplate, Spacer,
                                        Table, TableStyle)
    except ImportError as exc:  # noqa: F841
        raise RuntimeError(
            "O PDF precisa da biblioteca 'reportlab'. Instale uma vez na venv:\n"
            "    cd backend && venv\\Scripts\\activate && pip install reportlab\n"
            "(o Excel funciona sem ela). Depois tente exportar em PDF de novo."
        )

    r = resumir(chamados)
    gerado_em = datetime.now().strftime("%d/%m/%Y %H:%M")
    azul = colors.HexColor("#003366")

    styles = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=styles["Title"], textColor=azul, fontSize=18, spaceAfter=2)
    sub = ParagraphStyle("sub", parent=styles["Normal"], textColor=colors.grey, fontSize=9)
    h2 = ParagraphStyle("h2", parent=styles["Heading2"], textColor=azul, fontSize=12, spaceBefore=12)

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, title="Relatório de Chamados",
                            leftMargin=1.6 * cm, rightMargin=1.6 * cm,
                            topMargin=1.4 * cm, bottomMargin=1.4 * cm)
    elems = []

    def tabela(colunas, linhas, larguras, aligns=None):
        dados = [colunas] + linhas
        t = Table(dados, colWidths=larguras, repeatRows=1)
        est = [
            ("BACKGROUND", (0, 0), (-1, 0), azul),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F6FB")]),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D9D9D9")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]
        for col in (aligns or []):
            est.append(("ALIGN", (col, 1), (col, -1), "RIGHT"))
        t.setStyle(TableStyle(est))
        return t

    elems.append(Paragraph("Relatório de Chamados — Softdesk", h1))
    elems.append(Paragraph("Martinelli Advogados", sub))
    elems.append(Paragraph(f"Gerado em {gerado_em}", sub))
    elems.append(Paragraph(f"Filtro: {descricao_filtro or 'todos os chamados'}", sub))
    elems.append(Spacer(1, 10))

    # KPIs
    kpi_tab = tabela(
        ["Total", "Concluídos", "Em aberto", "Categoria mais frequente"],
        [[str(r["total"]), str(r["concluidos"]), str(r["abertos"]), r["categoria_top"]]],
        [3 * cm, 3 * cm, 3 * cm, 7.4 * cm],
    )
    elems.append(kpi_tab)

    def pct(v):
        return f"{round(v / r['total'] * 100) if r['total'] else 0}%"

    elems.append(Paragraph("Por categoria", h2))
    elems.append(tabela(["Categoria", "Qtde", "%"],
                        [[k, str(v), pct(v)] for k, v in r["por_categoria"]],
                        [11 * cm, 2.7 * cm, 2.7 * cm], aligns=[1, 2]))

    elems.append(Paragraph("Por status", h2))
    elems.append(tabela(["Status", "Qtde", "%"],
                        [[k, str(v), pct(v)] for k, v in r["por_status"]],
                        [11 * cm, 2.7 * cm, 2.7 * cm], aligns=[1, 2]))

    elems.append(Paragraph("Por mês", h2))
    elems.append(tabela(["Mês", "Abertos", "Concluídos", "% concluído"],
                        [[_fmt_mes(d["mes"]), str(d["total"]), str(d["concluidos"]), f"{d['pct']}%"]
                         for d in r["por_mes"]],
                        [4 * cm, 4 * cm, 4 * cm, 4.4 * cm], aligns=[1, 2, 3]))

    top = r["por_solicitante"][:15]
    elems.append(Paragraph("Top solicitantes", h2))
    elems.append(tabela(["Solicitante", "Qtde"],
                        [[k, str(v)] for k, v in top],
                        [13.4 * cm, 3 * cm], aligns=[1]))

    doc.build(elems)
    return buf.getvalue()
